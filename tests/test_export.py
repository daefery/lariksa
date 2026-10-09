from docx import Document

from lariksa.export import full_markdown, write_docx
from lariksa.notulen import Notulen
from lariksa.store import Meeting
from lariksa.transcript import ME, OTHERS, Utterance


def meeting(with_notulen=True):
    return Meeting(
        id="x", title="Rapat", created_at="2026-10-09T10:00:00", status="done",
        transcript=[Utterance(ME, 0, 2, "Halo"), Utterance(OTHERS, 3, 5, "Pagi")],
        notulen=Notulen.from_dict({
            "title": "Kickoff", "summary": "Mulai.", "decisions": ["Go"],
            "action_items": [{"task": "Siapkan deck", "owner": "Sari", "due": "Senin"}],
        }) if with_notulen else None,
    )


def test_docx_contains_notulen_table_and_transcript(tmp_path):
    target = tmp_path / "n.docx"
    write_docx(meeting(), target, date="9 Okt 2026")
    doc = Document(str(target))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Kickoff" in text and "Mulai." in text and "[00:03] Peserta: Pagi" in text
    cells = [c.text for c in doc.tables[0].rows[1].cells]
    assert cells == ["Siapkan deck", "Sari", "Senin"]


def test_english_docx_has_english_headings(tmp_path):
    m = meeting()
    m.notulen.language = "en"
    target = tmp_path / "en.docx"
    write_docx(m, target)
    doc = Document(str(target))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Summary" in text and "Transcript" in text and "Ringkasan" not in text
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Task", "Owner", "Due"]


def test_markdown_includes_transcript_even_without_notulen():
    md = full_markdown(meeting(with_notulen=False), date="9 Okt")
    assert md.startswith("# Rapat")
    assert "**[00:00] Saya:** Halo" in md
