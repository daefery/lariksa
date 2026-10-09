"""Exports: DOCX for the office, Markdown with the full transcript.

Headings follow the notulen's language (see ``notulen.LABELS``).
"""

from __future__ import annotations

from pathlib import Path

from .notulen import labels, to_markdown
from .store import Meeting
from .transcript import clock

_FOREST = (0x23, 0x36, 0x2C)
_MUTED = (0x65, 0x73, 0x6B)


def _language(meeting: Meeting) -> str:
    return meeting.notulen.language if meeting.notulen else "id"


def full_markdown(meeting: Meeting, *, date: str = "") -> str:
    t = labels(_language(meeting))
    parts = []
    if meeting.notulen:
        parts.append(to_markdown(meeting.notulen, date=date))
    else:
        parts.append(f"# {meeting.title}\n_{date}_\n")
    if meeting.transcript:
        parts.append(f"## {t['transcript']}\n")
        parts += [f"**[{clock(u.start)}] {u.speaker}:** {u.text}  " for u in meeting.transcript]
    return "\n".join(parts) + "\n"


def write_docx(meeting: Meeting, target: Path, *, date: str = "", with_transcript: bool = True) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    n = meeting.notulen
    t = labels(_language(meeting))
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def heading(text: str, level: int) -> None:
        h = doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.color.rgb = RGBColor(*_FOREST)

    def bullets(items: list[str]) -> None:
        for item in items or [t["none"]]:
            doc.add_paragraph(item, style="List Bullet")

    heading(n.title or meeting.title, 0)
    if date:
        meta = doc.add_paragraph(date)
        meta.runs[0].font.color.rgb = RGBColor(*_MUTED)

    heading(t["summary"], 1)
    doc.add_paragraph(n.summary or "-")
    heading(t["topics"], 1)
    bullets(n.topics)
    heading(t["decisions"], 1)
    bullets(n.decisions)

    heading(t["action_items"], 1)
    if n.action_items:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        for cell, text in zip(table.rows[0].cells, (t["task"], t["owner"], t["due"]), strict=True):
            cell.text = text
        for item in n.action_items:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = item.task, item.owner or "-", item.due or "-"
    else:
        doc.add_paragraph(t["none"])

    heading(t["open_questions"], 1)
    bullets(n.open_questions)

    if with_transcript and meeting.transcript:
        doc.add_page_break()
        heading(t["transcript"], 1)
        for u in meeting.transcript:
            p = doc.add_paragraph()
            label = p.add_run(f"[{clock(u.start)}] {u.speaker}: ")
            label.bold = True
            label.font.color.rgb = RGBColor(*_MUTED)
            p.add_run(u.text)

    doc.save(str(target))
